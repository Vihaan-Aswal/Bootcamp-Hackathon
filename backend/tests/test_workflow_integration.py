"""Back-to-back API integration for the five candidate-facing modules."""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.config import settings
from app.contracts.v1 import RADIX_CATEGORY_CODES


def _docx(heading: str, paragraph: str) -> bytes:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_two_roles_flow_through_all_modules_and_distinguish_candidates(
    client: TestClient, fixture_json, monkeypatch
) -> None:
    monkeypatch.setattr(settings, "llm_provider", "mock")

    health = client.get("/api/v1/health").json()
    assert {module["key"] for module in health["modules"]} == {
        "jd_analytics",
        "resume_parsing",
        "profile_builder",
        "talent_check",
        "skill_matching",
    }
    companies = client.get("/api/v1/companies").json()
    assert companies["companies"]
    company_id = companies["companies"][0]["company_id"]

    resume_sources = [
        (
            "candidate-alpha.docx",
            "Candidate Alpha",
            "I built a reporting app with Python and SQL.",
            10,
        ),
        (
            "candidate-beta.docx",
            "Candidate Beta",
            "I built a web app with JavaScript and AWS.",
            1,
        ),
    ]
    profiles: list[dict[str, object]] = []
    for index, (filename, name, resume_text, rating) in enumerate(resume_sources):
        parsed = client.post(
            "/api/v1/resumes/parse",
            files={
                "file": (
                    filename,
                    _docx("Projects", resume_text),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert parsed.status_code == 200, parsed.text
        resume = parsed.json()
        assert resume["source_type"] == "resume"
        assert resume["skills"]

        profile = fixture_json("candidate_profile.json")
        profile["profile_id"] = str(uuid4())
        profile["basic_details"] = {
            "full_name": name,
            "email": f"candidate-{index + 1}@example.com",
        }
        profile["skills"] = resume["skills"]
        profile["education"] = resume["education"]
        profile["projects"] = resume["projects"]
        profile["experience"] = resume["experience"]
        profile["cv_filename"] = resume["filename"]
        # These are explicit demo self-reports; no value is derived from resume text.
        profile["skillset_levels"] = [
            {"category_code": category.value, "level": rating, "source": "self_reported"}
            for category in RADIX_CATEGORY_CODES
        ]
        saved = client.post("/api/v1/profiles/validate", json=profile)
        assert saved.status_code == 200, saved.text
        assert "raw_resume_text" not in saved.json()
        profiles.append(saved.json())

    role_sources = [
        ("python-sql-role.docx", "Required skills: Python and SQL."),
        ("javascript-cloud-role.docx", "Required skills: JavaScript and AWS."),
    ]
    outcomes: dict[tuple[int, int], int] = {}
    readiness_scores: list[int] = []

    for role_index, (filename, description) in enumerate(role_sources):
        analyzed = client.post(
            "/api/v1/jds/analyze",
            files={
                "file": (
                    filename,
                    _docx("Requirements", description),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert analyzed.status_code == 200, analyzed.text
        jd = analyzed.json()
        assert jd["source_type"] == "jd"
        assert jd["skills"]

        for candidate_index, profile in enumerate(profiles):
            checked = client.post(
                f"/api/v1/companies/{company_id}/talent-check", json=profile
            )
            assert checked.status_code == 200, checked.text
            talent_check = checked.json()
            assert talent_check["status"] == "complete"
            assert len(talent_check["rows"]) == 12
            readiness_scores.append(talent_check["readiness_score"])

            matched = client.post(
                "/api/v1/skill-matches/score",
                json={
                    "schema_version": "1.0.0",
                    "profile": profile,
                    "jd": jd,
                    "confirmations": [],
                },
            )
            assert matched.status_code == 200, matched.text
            result = matched.json()
            assert result["status"] == "scorable"
            outcomes[(role_index, candidate_index)] = result["match_score"]

    assert outcomes[(0, 0)] == 100
    assert outcomes[(0, 1)] == 0
    assert outcomes[(1, 0)] == 0
    assert outcomes[(1, 1)] == 100
    assert len(set(readiness_scores)) > 1
