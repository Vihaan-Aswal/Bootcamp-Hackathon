"""Offline tests for private, evidence-backed resume parsing."""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient

from app.api.dependencies import require_document_provider
from app.api.routes.resumes import parse_resume_endpoint, router
from app.contracts.v1 import SourceElement
from app.parsing.document import DocumentParseError, parse_document
from app.parsing.privacy import redact_resume_identifiers
from app.providers.base import ProviderError, StructuredGenerationRequest
from app.providers.mock import MockProvider
from app.services import resume_parsing
from app.services.resume_parsing import parse_resume


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Avery Candidate", level=1)
    document.add_paragraph("Email: avery@example.test | Phone: +1 415 555 2671")
    document.add_heading("Projects", level=1)
    document.add_paragraph("Built a reporting dashboard in Python and SQL.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Course"
    table.cell(0, 1).text = "Applied Data Systems"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes(text: str = "Built a reporting dashboard in Python and SQL.") -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        ),
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode("ascii"))
        content.extend(body + b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
        ).encode("ascii")
    )
    return bytes(content)


def _client() -> TestClient:
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[require_document_provider] = MockProvider
    return TestClient(app)


def test_parser_preserves_docx_paragraph_and_table_order() -> None:
    parsed = parse_document("resume.docx", _docx_bytes(), max_bytes=1_000_000, max_pages=10)

    assert parsed.source_format == "docx"
    assert [element.order for element in parsed.elements] == list(range(len(parsed.elements)))
    assert any(
        element.kind.value == "heading" and element.text == "Projects"
        for element in parsed.elements
    )
    table_row = next(element for element in parsed.elements if element.kind.value == "table_row")
    assert "Applied Data Systems" in table_row.text
    assert table_row.paragraph_index is not None


def test_pdf_parse_returns_page_location_and_offline_skill_suggestion() -> None:
    parsed = parse_document("resume.pdf", _pdf_bytes(), max_bytes=1_000_000, max_pages=10)
    result = parse_resume(
        filename=parsed.filename,
        content=_pdf_bytes(),
        provider=MockProvider(),
    )

    assert parsed.page_count == 1
    assert result.source_type.value == "resume"
    assert {skill.skill_id for skill in result.skills} >= {"python", "sql"}
    assert all(skill.source_ref and skill.source_ref.page == 1 for skill in result.skills)
    assert not hasattr(result, "skillset_levels")


def test_resume_endpoint_returns_shared_contract_and_closes_upload() -> None:
    upload = UploadFile(filename="resume.docx", file=BytesIO(_docx_bytes()))

    result = parse_resume_endpoint(file=upload, provider=MockProvider())

    assert result.filename == "resume.docx"
    assert any(skill.skill_id == "python" for skill in result.skills)
    assert upload.file.closed


def test_resume_api_accepts_multipart_pdf() -> None:
    with _client() as client:
        response = client.post(
            "/api/v1/resumes/parse",
            files={"file": ("resume.pdf", _pdf_bytes(), "application/pdf")},
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["source_type"] == "resume"
    assert any(skill["skill_id"] == "python" for skill in payload["skills"])


def test_remote_processing_requires_backend_opt_in_before_provider_call(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(resume_parsing.settings, "llm_resume_processing_allowed", False)
    calls: list[StructuredGenerationRequest] = []

    class RemoteProvider:
        name = "test-provider"
        model = "test-model"
        supports_json_schema = True

        def generate_structured(self, request: StructuredGenerationRequest) -> str:
            calls.append(request)
            return '{"skills":[],"education":[],"projects":[],"experience":[]}'

    with pytest.raises(ProviderError, match="current terms"):
        parse_resume(
            filename="resume.docx",
            content=_docx_bytes(),
            provider=RemoteProvider(),
            allow_external_processing=True,
        )

    assert calls == []


def test_remote_processing_requires_per_request_consent(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(resume_parsing.settings, "llm_resume_processing_allowed", True)

    class RemoteProvider:
        name = "test-provider"
        model = "test-model"
        supports_json_schema = True

        def generate_structured(self, request: StructuredGenerationRequest) -> str:
            pytest.fail("provider must not be called without user consent")

    with pytest.raises(ProviderError, match="Confirm that you agree"):
        parse_resume(
            filename="resume.docx",
            content=_docx_bytes(),
            provider=RemoteProvider(),
            allow_external_processing=False,
        )


def test_external_request_redacts_identifiers_and_verifies_against_original(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(resume_parsing.settings, "llm_resume_processing_allowed", True)
    seen: list[StructuredGenerationRequest] = []

    class RemoteProvider:
        name = "test-provider"
        model = "test-model"
        supports_json_schema = True

        def generate_structured(self, request: StructuredGenerationRequest) -> str:
            seen.append(request)
            assert "avery@example.test" not in request.user_prompt.casefold()
            assert "415 555 2671" not in request.user_prompt
            assert "Avery Candidate" not in request.user_prompt
            assert all(
                "Avery Candidate" not in (element.heading or "") for element in request.elements
            )
            project_element = next(
                element for element in request.elements if "Python and SQL" in element.text
            )
            return json.dumps(
                {
                    "skills": [
                        {
                            "skill_name": "Python",
                            "category_code": "COD",
                            "requirement_type": "unclear",
                            "confidence": "high",
                            "evidence_quote": "Built a reporting dashboard in Python and SQL.",
                            "source_element_id": project_element.element_id,
                        }
                    ],
                    "education": [],
                    "projects": [],
                    "experience": [],
                }
            )

    result = parse_resume(
        filename="resume.docx",
        content=_docx_bytes(),
        provider=RemoteProvider(),
        allow_external_processing=True,
    )

    assert len(seen) == 1
    assert result.metadata.redaction_applied is True
    assert result.skills[0].evidence_status.value == "exact"
    assert result.skills[0].source_ref and result.skills[0].source_ref.paragraph is not None


def test_direct_identifier_redaction_masks_name_email_and_phone() -> None:
    source = (
        SourceElement(element_id="d0", text="Avery Candidate", order=0),
        SourceElement(
            element_id="d1",
            text="Email: avery@example.test | Phone: +1 415 555 2671",
            order=1,
        ),
    )

    redacted, applied = redact_resume_identifiers(source)

    assert applied is True
    assert all("Avery Candidate" not in element.text for element in redacted)
    assert "avery@example.test" not in " ".join(element.text for element in redacted)
    assert "415 555 2671" not in " ".join(element.text for element in redacted)


def test_scanned_pdf_returns_ocr_required_without_provider_call() -> None:
    scanned = _pdf_bytes(" ")

    with pytest.raises(DocumentParseError) as error:
        parse_resume(filename="resume.pdf", content=scanned, provider=MockProvider())

    assert error.value.code == "OCR_REQUIRED"


def test_endpoint_returns_clear_unsupported_format_error() -> None:
    upload = UploadFile(filename="resume.txt", file=BytesIO(b"plain text"))

    with pytest.raises(HTTPException) as error:
        parse_resume_endpoint(file=upload, provider=MockProvider())

    assert getattr(error.value, "status_code", None) == 415
    assert upload.file.closed
