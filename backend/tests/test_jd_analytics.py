"""Synthetic offline coverage for JD Analytics and upload lifecycle."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from starlette.datastructures import UploadFile

from app.api.dependencies import require_document_provider
from app.api.routes.jds import analyze_jd_upload, router
from app.config import settings
from app.contracts.dto import LlmExtractionOutput, LlmSkillClaim
from app.contracts.v1 import RequirementType, SourceElement
from app.parsing.document import parse_document
from app.providers.base import StructuredGenerationRequest
from app.providers.mock import MockProvider
from app.services.jd_analytics import analyze_jd, classify_requirement


def _make_pdf(lines: list[str]) -> bytes:
    def pdf_escape(value: str) -> bytes:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode()

    commands = [b"BT", b"/F1 11 Tf"]
    for row, line in enumerate(lines):
        commands.extend(
            (
                f"1 0 0 1 72 {740 - row * 18} Tm".encode(),
                b"(" + pdf_escape(line) + b") Tj",
            )
        )
    commands.append(b"ET")
    stream = b"\n".join(commands)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%synthetic\n")
    offsets = [0]
    for number, content in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + content + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def _make_docx() -> bytes:
    document = Document()
    document.add_heading("What We're Looking For", level=1)
    document.add_paragraph("Python is required for this software engineering role.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Preferred technology"
    table.rows[0].cells[1].text = "Docker is a plus for this team."
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _client(provider: MockProvider) -> TestClient:
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[require_document_provider] = lambda: provider
    return TestClient(app)


def test_jd_pdf_returns_contract_with_verified_locations() -> None:
    content = _make_pdf(
        [
            "Key Responsibilities",
            "Python is required, and Docker is a plus for cloud deployment.",
            "You will design REST APIs and review code.",
        ]
    )

    with _client(MockProvider()) as client:
        response = client.post("/api/v1/jds/analyze", files={"file": ("role.pdf", content)})

    assert response.status_code == 200, response.text
    result = response.json()
    by_name = {skill["skill_name"]: skill for skill in result["skills"]}
    assert by_name["Python"]["requirement_type"] == "required"
    assert by_name["Docker"]["requirement_type"] == "preferred"
    assert by_name["REST APIs"]["requirement_type"] == "mentioned"
    assert by_name["Python"]["evidence_status"] == "exact"
    assert by_name["Python"]["source_ref"]["page"] == 1
    assert by_name["Python"]["source_ref"]["element_id"].startswith("pdf-p1-")


def test_jd_docx_preserves_heading_paragraph_and_table_order() -> None:
    document = parse_document(
        "role.docx",
        _make_docx(),
        max_bytes=1_000_000,
        max_pages=20,
    )

    assert document.source_format == "docx"
    assert [element.kind.value for element in document.elements] == [
        "heading",
        "paragraph",
        "table_row",
    ]
    assert [element.order for element in document.elements] == [0, 1, 2]
    assert document.elements[1].heading == "What We're Looking For"
    assert document.elements[2].paragraph_index is not None


def test_requirement_label_requires_matching_source_language() -> None:
    element = parse_document("role.docx", _make_docx(), max_bytes=1_000_000, max_pages=20).elements[
        1
    ]
    claim = LlmSkillClaim(
        skill_name="Python",
        evidence_quote=element.text,
        source_element_id=element.element_id,
        requirement_type="required",
    )

    assert classify_requirement(claim, element) is RequirementType.required

    claim = claim.model_copy(update={"evidence_quote": "Python is a bonus for this role."})
    assert classify_requirement(claim, element) is RequirementType.preferred

    claim = claim.model_copy(update={"evidence_quote": "Python is not required."})
    assert classify_requirement(claim, element) is RequirementType.unclear

    claim = claim.model_copy(update={"evidence_quote": "Experience with Python."})
    assert classify_requirement(claim, element) is RequirementType.mentioned

    claim = claim.model_copy(
        update={
            "evidence_quote": "Python is useful for the team.",
            "requirement_type": RequirementType.required,
        }
    )
    assert classify_requirement(claim, element) is RequirementType.mentioned

    claim = claim.model_copy(update={"evidence_quote": "You must not use Python here."})
    assert classify_requirement(claim, element) is RequirementType.unclear


def test_trailing_requirement_cue_applies_to_every_skill_in_a_list() -> None:
    evidence = "Python, SQL and Java are required for this role."
    element = SourceElement(
        element_id="p1",
        text=evidence,
        kind="paragraph",
        order=0,
    )

    for skill_name in ("Python", "SQL", "Java"):
        claim = LlmSkillClaim(
            skill_name=skill_name,
            evidence_quote=evidence,
            source_element_id=element.element_id,
            requirement_type="mentioned",
        )
        assert classify_requirement(claim, element) is RequirementType.required


def test_separate_cues_keep_each_skill_requirement_distinct() -> None:
    evidence = "Python is required, while Docker is a plus for the role."
    element = SourceElement(
        element_id="p2",
        text=evidence,
        kind="paragraph",
        order=0,
    )
    python_claim = LlmSkillClaim(
        skill_name="Python",
        evidence_quote=evidence,
        source_element_id=element.element_id,
        requirement_type="preferred",
    )
    docker_claim = LlmSkillClaim(
        skill_name="Docker",
        evidence_quote=evidence,
        source_element_id=element.element_id,
        requirement_type="required",
    )

    assert classify_requirement(python_claim, element) is RequirementType.required
    assert classify_requirement(docker_claim, element) is RequirementType.preferred


def test_unverified_provider_claim_is_rejected_not_returned_as_a_skill() -> None:
    content = _make_pdf(["Requirements", "Python is essential for this position."])

    def fabricated_quote(_request: StructuredGenerationRequest) -> str:
        element = parse_document("role.pdf", content, max_bytes=1_000_000, max_pages=20).elements[1]
        return LlmExtractionOutput(
            skills=[
                LlmSkillClaim(
                    skill_name="Python",
                    evidence_quote="Expert in Rust",
                    source_element_id=element.element_id,
                    requirement_type="required",
                )
            ]
        ).model_dump_json()

    result = analyze_jd("role.pdf", content, MockProvider(fabricated_quote))

    assert result.skills == []
    assert result.rejected_claims[0].claim == "Python"
    assert "not found" in result.rejected_claims[0].reason


def test_bad_and_oversized_uploads_return_safe_errors() -> None:
    with _client(MockProvider()) as client:
        bad = client.post("/api/v1/jds/analyze", files={"file": ("role.txt", b"not a JD")})
        oversized = b"%PDF-" + b"x" * (settings.max_upload_mb * 1024 * 1024)
        huge = client.post("/api/v1/jds/analyze", files={"file": ("role.pdf", oversized)})

    assert bad.status_code == 415
    assert bad.json()["detail"]["code"] == "UNSUPPORTED_FORMAT"
    assert huge.status_code == 413
    assert huge.json()["detail"]["code"] == "FILE_TOO_LARGE"


def test_scanned_pdf_reports_ocr_required() -> None:
    with _client(MockProvider()) as client:
        response = client.post(
            "/api/v1/jds/analyze",
            files={"file": ("scan.pdf", _make_pdf([]))},
        )

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "OCR_REQUIRED"
    assert "OCR" in response.json()["detail"]["message"]


def test_malformed_provider_response_gets_one_bounded_retry_then_safe_error() -> None:
    calls = 0

    def malformed(_request: StructuredGenerationRequest) -> str:
        nonlocal calls
        calls += 1
        return "{truncated"

    content = _make_pdf(["Qualifications", "Python is required for this position."])
    with _client(MockProvider(malformed)) as client:
        response = client.post("/api/v1/jds/analyze", files={"file": ("role.pdf", content)})

    assert calls == 2
    assert response.status_code == 502
    assert response.json()["detail"]["code"] == "INVALID_PROVIDER_OUTPUT"


def test_upload_stream_is_closed_on_success_and_parse_failure() -> None:
    class CloseSpy(BytesIO):
        close_called = False

        def close(self) -> None:
            self.close_called = True
            super().close()

    success_stream = CloseSpy(_make_pdf(["Requirements", "Python is required for this role."]))
    success_upload = UploadFile(filename="role.pdf", file=success_stream)
    analyze_jd_upload(success_upload, MockProvider())
    assert success_stream.close_called

    bad_stream = CloseSpy(b"invalid")
    bad_upload = UploadFile(filename="role.pdf", file=bad_stream)
    try:
        analyze_jd_upload(bad_upload, MockProvider())
    except Exception:
        pass
    assert bad_stream.close_called


def test_upload_read_oserror_returns_safe_400_and_closes_stream() -> None:
    class ReadFailure(BytesIO):
        close_called = False

        def read(self, _size: int = -1) -> bytes:
            raise OSError("private OS details should not be returned")

        def close(self) -> None:
            self.close_called = True
            super().close()

    stream = ReadFailure(b"ignored")
    upload = UploadFile(filename="role.pdf", file=stream)

    try:
        analyze_jd_upload(upload, MockProvider())
    except HTTPException as error:
        assert error.status_code == 400
        assert error.detail == {
            "code": "UPLOAD_READ_FAILED",
            "message": "The uploaded file could not be read.",
        }
    else:
        raise AssertionError("expected a safe HTTP 400 response")

    assert stream.close_called
