"""Build labelled DOCX inputs for the no-network UI demo."""

from __future__ import annotations

from pathlib import Path

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]
SYNTHETIC_ROOT = REPO_ROOT / "data" / "samples" / "synthetic"

FIXTURES = {
    "JDs/role_alpha.docx": (
        "Synthetic Demo Role Alpha",
        "Required Qualifications",
        "Python and SQL are required for this demo role.",
    ),
    "JDs/role_beta.docx": (
        "Synthetic Demo Role Beta",
        "Required Qualifications",
        "JavaScript and AWS are required for this demo role.",
    ),
    "Resumes/candidate_alpha.docx": (
        "Synthetic Candidate Alpha",
        "Projects",
        "Built a data reporting demo using Python and SQL.",
    ),
    "Resumes/candidate_beta.docx": (
        "Synthetic Candidate Beta",
        "Projects",
        "Built a cloud web demo using JavaScript and AWS.",
    ),
}


def main() -> None:
    for relative_path, (title, heading, paragraph) in FIXTURES.items():
        document = Document()
        document.add_heading(title, level=0)
        document.add_heading(heading, level=1)
        document.add_paragraph(paragraph)
        output = SYNTHETIC_ROOT / relative_path
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
    print("Built four labelled synthetic DOCX demo inputs.")


if __name__ == "__main__":
    main()
