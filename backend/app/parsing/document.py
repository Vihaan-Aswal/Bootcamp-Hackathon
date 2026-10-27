"""Safe, format-specific parsing behind the shared M2 document interface."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath, PurePosixPath
from typing import Literal

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import ValidationError

from app.contracts.v1 import ElementKind, SourceElement

DocumentFormat = Literal["pdf", "docx"]


class DocumentParseError(ValueError):
    """A safe, user-facing parse failure that never includes source contents."""

    def __init__(self, code: str, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    """The stable parser output consumed by JD and resume services.

    ``elements`` are ordered and carry parser-created IDs and locations. PDF
    elements have page numbers; DOCX elements have paragraph indexes where
    available. Model output never supplies or changes these locations.
    """

    filename: str
    source_format: DocumentFormat
    elements: tuple[SourceElement, ...]
    page_count: int | None
    character_count: int

    @property
    def text(self) -> str:
        return "\n".join(element.text for element in self.elements)


_HEADING_NAMES = {
    "about",
    "about the role",
    "about us",
    "additional information",
    "basic qualifications",
    "career summary",
    "certifications",
    "education",
    "experience",
    "job description",
    "key responsibilities",
    "minimum qualifications",
    "preferred qualifications",
    "projects",
    "qualifications",
    "requirements",
    "responsibilities",
    "skills",
    "technical skills",
    "what we are looking for",
    "what we're looking for",
    "what you will do",
    "what you'll do",
    "who you are",
    "who we are",
}
_BULLET_PREFIX = re.compile(r"^\s*(?:[•●▪◦‣*-]|\d{1,3}[.)])\s+")
_MAX_DOCX_ARCHIVE_ENTRIES = 2_000
_MAX_DOCX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
_MAX_DOCX_COMPRESSION_RATIO = 150


def safe_filename(filename: str | None, fallback: str = "upload") -> str:
    """Keep a display filename, never a client-supplied path or control chars."""

    leaf = (filename or fallback).replace("\\", "/").rsplit("/", 1)[-1]
    cleaned = "".join(character for character in leaf if character.isprintable())
    cleaned = cleaned.strip(" .")
    return (cleaned or fallback)[:255]


def _kind_for_line(text: str, *, style_name: str = "") -> ElementKind:
    normalized = re.sub(r"\s+", " ", text).strip().rstrip(":").casefold()
    if style_name.casefold().startswith("heading") or normalized in _HEADING_NAMES:
        return ElementKind.heading
    if style_name.casefold().startswith("list") or _BULLET_PREFIX.match(text):
        return ElementKind.list_item
    return ElementKind.paragraph


def _looks_unusable(text: str) -> bool:
    if len(text.strip()) < 20:
        return True
    printable = sum(character.isprintable() or character.isspace() for character in text)
    alphanumeric = sum(character.isalnum() for character in text)
    return printable / max(1, len(text)) < 0.75 or alphanumeric < 12


def parse_document(
    filename: str | None,
    content: bytes,
    *,
    max_bytes: int,
    max_pages: int,
    max_characters: int = 120_000,
) -> ParsedDocument:
    """Validate and parse one PDF/DOCX from memory, with no persistent copy.

    Raises ``DocumentParseError`` with stable machine-readable ``code`` values
    and safe messages. The caller must close its UploadFile in a ``finally``
    block; this parser does not retain uploaded data.
    """

    display_name = safe_filename(filename)
    suffix = PurePath(display_name).suffix.casefold()
    if suffix not in {".pdf", ".docx"}:
        raise DocumentParseError(
            "UNSUPPORTED_FORMAT", "Upload a PDF or DOCX file.", status_code=415
        )
    if not content:
        raise DocumentParseError("EMPTY_FILE", "The uploaded file is empty.")
    if len(content) > max_bytes:
        raise DocumentParseError(
            "FILE_TOO_LARGE", f"The file exceeds the {max_bytes // (1024 * 1024)} MB limit.", 413
        )

    if suffix == ".pdf":
        if b"%PDF-" not in content[:1024]:
            raise DocumentParseError("INVALID_SIGNATURE", "The file is not a valid PDF.")
        return _parse_pdf(display_name, content, max_pages=max_pages, max_characters=max_characters)

    if content.startswith(b"\xd0\xcf\x11\xe0"):
        raise DocumentParseError(
            "ENCRYPTED_FILE", "Password-protected Office files are not supported."
        )
    if not content.startswith(b"PK\x03\x04"):
        raise DocumentParseError("INVALID_SIGNATURE", "The file is not a valid DOCX document.")
    return _parse_docx(display_name, content, max_characters=max_characters)


def _parse_pdf(
    filename: str,
    content: bytes,
    *,
    max_pages: int,
    max_characters: int,
) -> ParsedDocument:
    try:
        with pdfplumber.open(io.BytesIO(content), password="") as pdf:
            page_count = len(pdf.pages)
            if page_count == 0:
                raise DocumentParseError("EMPTY_DOCUMENT", "The PDF has no pages.")
            if page_count > max_pages:
                raise DocumentParseError(
                    "TOO_MANY_PAGES", f"The PDF exceeds the {max_pages} page limit.", 413
                )

            elements: list[SourceElement] = []
            order = 0
            heading: str | None = None
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                for line_number, raw_line in enumerate(text.splitlines(), start=1):
                    line = re.sub(r"\s+", " ", raw_line).strip()
                    if not line:
                        continue
                    kind = _kind_for_line(line)
                    if kind is ElementKind.heading:
                        heading = line.rstrip(":")[:200]
                    elements.append(
                        SourceElement(
                            element_id=f"pdf-p{page_number}-l{line_number}",
                            text=line,
                            kind=kind,
                            order=order,
                            heading=heading,
                            page=page_number,
                        )
                    )
                    order += 1
    except DocumentParseError:
        raise
    except Exception as exc:
        if (
            "password" in type(exc).__name__.casefold()
            or "encrypt" in type(exc).__name__.casefold()
        ):
            raise DocumentParseError(
                "ENCRYPTED_FILE", "Password-protected PDFs are not supported."
            ) from None
        raise DocumentParseError(
            "INVALID_DOCUMENT", "The PDF could not be read. It may be corrupt or encrypted."
        ) from None

    character_count = sum(len(element.text) for element in elements)
    if character_count > max_characters:
        raise DocumentParseError(
            "DOCUMENT_TOO_LONG", "The extracted text exceeds the processing limit.", 413
        )
    if not elements or _looks_unusable("\n".join(element.text for element in elements)):
        raise DocumentParseError(
            "OCR_REQUIRED",
            "No usable text was found. This PDF may be scanned; OCR is required before it can "
            "be analyzed.",
        )
    return ParsedDocument(filename, "pdf", tuple(elements), page_count, character_count)


def _validate_docx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > _MAX_DOCX_ARCHIVE_ENTRIES:
                raise DocumentParseError(
                    "INVALID_DOCUMENT", "The DOCX archive has too many entries."
                )
            if any(".." in PurePosixPath(entry.filename).parts for entry in entries):
                raise DocumentParseError(
                    "INVALID_DOCUMENT", "The DOCX archive contains an unsafe path."
                )
            if sum(entry.file_size for entry in entries) > _MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentParseError(
                    "INVALID_DOCUMENT", "The DOCX archive is too large to unpack safely."
                )
            if any(
                entry.file_size > 0
                and entry.file_size / max(1, entry.compress_size) > _MAX_DOCX_COMPRESSION_RATIO
                for entry in entries
            ):
                raise DocumentParseError(
                    "INVALID_DOCUMENT", "The DOCX archive compression ratio is unsafe."
                )
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentParseError(
                    "INVALID_DOCUMENT", "The DOCX archive is missing required document data."
                )
            if "word/vbaProject.bin" in names:
                raise DocumentParseError(
                    "UNSUPPORTED_FORMAT", "Macro-enabled Word files are not supported.", 415
                )
            if archive.testzip() is not None:
                raise DocumentParseError("INVALID_DOCUMENT", "The DOCX archive is corrupt.")
    except DocumentParseError:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError):
        raise DocumentParseError(
            "INVALID_DOCUMENT", "The DOCX file is corrupt or unreadable."
        ) from None


def _parse_docx(filename: str, content: bytes, *, max_characters: int) -> ParsedDocument:
    _validate_docx_archive(content)
    try:
        document = Document(io.BytesIO(content))
        body = document.element.body
        elements: list[SourceElement] = []
        paragraph_index = 0
        order = 0
        heading: str | None = None

        def add_element(text: str, kind: ElementKind, paragraph: int | None) -> None:
            nonlocal order, heading
            clean = re.sub(r"\s+", " ", text).strip()
            if not clean:
                return
            if kind is ElementKind.heading:
                heading = clean.rstrip(":")[:200]
            elements.append(
                SourceElement(
                    element_id=f"docx-{order}",
                    text=clean,
                    kind=kind,
                    order=order,
                    heading=heading,
                    paragraph_index=paragraph,
                )
            )
            order += 1

        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, document)
                current_index = paragraph_index
                paragraph_index += 1
                text = paragraph.text.strip()
                if text:
                    add_element(
                        text,
                        _kind_for_line(text, style_name=paragraph.style.name or ""),
                        current_index,
                    )
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                for row in table.rows:
                    cells: list[str] = []
                    first_index: int | None = None
                    for cell in row.cells:
                        for _paragraph in cell.paragraphs:
                            if first_index is None:
                                first_index = paragraph_index
                            paragraph_index += 1
                        text = re.sub(r"\s+", " ", cell.text).strip()
                        if text:
                            cells.append(text)
                    if cells:
                        add_element(" | ".join(cells), ElementKind.table_row, first_index)
    except DocumentParseError:
        raise
    except Exception:
        raise DocumentParseError("INVALID_DOCUMENT", "The DOCX file could not be read.") from None

    character_count = sum(len(element.text) for element in elements)
    if character_count > max_characters:
        raise DocumentParseError(
            "DOCUMENT_TOO_LONG", "The extracted text exceeds the processing limit.", 413
        )
    if not elements or _looks_unusable("\n".join(element.text for element in elements)):
        raise DocumentParseError("EMPTY_DOCUMENT", "The DOCX file contains no usable text.")
    try:
        # Re-validate parser-created values at the shared contract boundary.
        checked = tuple(SourceElement.model_validate(element.model_dump()) for element in elements)
    except ValidationError:
        raise DocumentParseError(
            "INVALID_DOCUMENT", "The DOCX contains invalid text structure."
        ) from None
    return ParsedDocument(filename, "docx", checked, None, character_count)
